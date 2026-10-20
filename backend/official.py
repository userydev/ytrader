"""Official exchange snapshots, independently bounded; Eastmoney is a labelled fallback."""
from __future__ import annotations

import io
import json
import re
import zipfile

from openpyxl import load_workbook

from .adapters import PublicData, UpstreamError
from .models import utcnow

SSE = 'https://query.sse.com.cn/sseQuery/commonQuery.do'
SZSE = 'https://www.szse.cn/api/report/ShowReport'
BSE = 'https://www.bse.cn/nqxxController/nqxxCnzq.do'


def security(exchange, code, name, industry, source, url):
    code = str(code).split('.')[0].zfill(6)
    if not re.fullmatch(r'\d{6}', code) or not str(name).strip():
        raise UpstreamError('官方证券条目缺少代码或名称')
    industry = str(industry).strip() if industry is not None else None
    if industry in ('', '-', 'None', '不适用'):
        industry = None
    return dict(symbol=exchange + ':' + code, code=code, name=str(name).strip(), exchange=exchange,
                industry=industry, status='listed', as_of=None, source=source, _source_url=url)


class ExchangeData(PublicData):
    def __init__(self, network):
        super().__init__(network)
        self.observations = {}
        self.reported = {}
        self.report_dates = {}

    def cross_total(self, exchange):
        filters = {'SH': 'm:1+t:2,m:1+t:23', 'SZ': 'm:0+t:6,m:0+t:80', 'BJ': 'm:0+t:81+s:2048'}
        try:
            payload = self.network.json('GET', 'https://push2.eastmoney.com/api/qt/clist/get', public=True,
                params={'pn': 1, 'pz': 1, 'np': 1, 'fs': filters[exchange], 'fields': 'f12'})
            total = (payload.get('data') or {}).get('total')
            return total if isinstance(total, int) and total > 0 else None
        except (UpstreamError, ValueError, TypeError):
            return None

    def securities(self, exchange):
        try:
            rows = getattr(self, exchange.lower())()
            minimum = {'SH': 1000, 'SZ': 1500, 'BJ': 150}[exchange]
            if len({r['symbol'] for r in rows}) != len(rows) or len(rows) < minimum:
                raise UpstreamError('官方证券快照少于覆盖安全阈值或含重复代码')
            cross = self.cross_total(exchange)
            self.observations[exchange] = dict(official_reported=self.reported.get(exchange), official_observed=len(rows),
                source_report_dates=self.report_dates.get(exchange, []), fetched_at=utcnow(),
                cross_source='东方财富公开分类列表（口径可能不同，不自动合并）', cross_reported=cross,
                reconciliation='unverified' if cross is None else ('consistent_counts_only' if cross == len(rows) else 'conflict'))
            return rows, cross == len(rows), self.reported.get(exchange, len(rows))
        except (UpstreamError, ValueError, KeyError, TypeError, zipfile.BadZipFile):
            rows, complete, total = super().securities(exchange)
            for row in rows:
                row['source'] = '东方财富公开证券列表（官方接口失败后的备用源）'
                row['_source_url'] = 'https://quote.eastmoney.com/center/gridlist.html'
            self.observations[exchange] = dict(official_reported=None, official_observed=None, cross_reported=total, reconciliation='official_unavailable')
            return rows, complete, total

    def sh(self):
        rows = []
        reported_total = 0
        for stock_type in ('1', '8'):
            expected = None
            part = []
            for page in range(1, 31):
                payload = self.network.json('GET', SSE, public=True,
                    headers={'Referer': 'https://www.sse.com.cn/assortment/stock/list/share/', 'User-Agent': 'Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/110.0.0.0 Safari/537.36'},
                    params={'sqlId': 'COMMON_SSE_CP_GPJCTPZ_GPLB_GP_L', 'STOCK_TYPE': stock_type,
                            'COMPANY_STATUS': '2,4,5,7,8', 'type': 'inParams', 'isPagination': 'true',
                            'REG_PROVINCE': '', 'CSRC_CODE': '', 'STOCK_CODE': '',
                            'jsonCallBack': 'jsonpCallback', 'pageHelp.pageSize': 10000, 'pageHelp.pageNo': page,
                            'pageHelp.beginPage': page, 'pageHelp.endPage': page, 'pageHelp.cacheSize': 1})
                result = payload.get('result')
                if not isinstance(result, list):
                    result = (payload.get('pageHelp') or {}).get('data')
                if not isinstance(result, list) or not result:
                    raise UpstreamError('上交所列表缺失')
                count = (payload.get('pageHelp') or {}).get('total')
                if count is not None:
                    count = int(count)
                    if expected is not None and count != expected:
                        raise UpstreamError('上交所分页总数变化')
                    expected = count
                for row in result:
                    part.append(security('SH', row.get('A_STOCK_CODE') or row.get('SEC_CODE'),
                                         row.get('COMPANY_ABBR') or row.get('SEC_NAME_CN'), row.get('CSRC_CODE_DESC'),
                                         '上海证券交易所公开 A 股列表（源发布时间未知）',
                                         'https://www.sse.com.cn/assortment/stock/list/share/'))
                if (expected is not None and len(part) >= expected) or (expected is None and len(result) < 10000):
                    break
            if expected is not None and len(part) != expected:
                raise UpstreamError('上交所分页覆盖不完整')
            if not part:
                raise UpstreamError('上交所板块为空')
            rows.extend(part)
            reported_total += expected if expected is not None else len(part)
        self.reported['SH'] = reported_total
        return rows

    def sz(self):
        raw = self.network.request('GET', SZSE, public=True,
            params={'SHOWTYPE': 'xlsx', 'CATALOGID': '1110', 'TABKEY': 'tab1'},
            headers={'Referer': 'https://www.szse.cn/', 'User-Agent': 'Mozilla/5.0'})
        with zipfile.ZipFile(io.BytesIO(raw)) as archive:
            if sum(i.file_size for i in archive.infolist()) > 40_000_000:
                raise UpstreamError('深交所工作簿解压大小超过限制')
        book = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
        try:
            sheet = book.active
            sheet.reset_dimensions()  # Official XLSX may incorrectly declare dimension A1.
            iterator = iter(sheet.iter_rows(values_only=True))
            headers = None
            for _ in range(12):
                row = next(iterator, ())
                labels = [str(v or '').strip() for v in row]
                if 'A股代码' in labels or '证券代码' in labels:
                    headers = labels
                    break
            if not headers:
                raise UpstreamError('深交所工作簿表头改变')
            rows = []
            for values in iterator:
                row = dict(zip(headers, values))
                code = row.get('A股代码') or row.get('证券代码')
                if not code:
                    continue
                rows.append(security('SZ', code, row.get('A股简称') or row.get('证券简称'), row.get('所属行业'),
                                     '深圳证券交易所公开 A 股列表（源发布时间未知）',
                                     SZSE + '?SHOWTYPE=xlsx&CATALOGID=1110&TABKEY=tab1'))
            self.reported['SZ'] = len(rows)
            return rows
        finally:
            book.close()

    def bj(self):
        rows, expected, report_dates = [], None, set()
        for page in range(100):
            raw = self.network.request('POST', BSE, public=True,
                form={'page': str(page), 'typejb': 'T', 'xxfcbj[]': '2', 'xxzqdm': '',
                      'sortfield': 'xxzqdm', 'sorttype': 'asc'},
                headers={'Referer': 'https://www.bse.cn/', 'User-Agent': 'Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/110.0.0.0 Safari/537.36'})
            decoded = raw.decode('utf-8')
            start, end = decoded.find('['), decoded.rfind(']')
            if start < 0 or end < start:
                raise UpstreamError('北交所列表无法解析')
            payload = json.loads(decoded[start:end + 1])
            data = payload[0]
            content = data.get('content')
            if not isinstance(content, list):
                raise UpstreamError('北交所 content 缺失')
            total = data.get('totalElements') or data.get('total')
            if total is not None:
                total = int(total)
                if expected is not None and expected != total:
                    raise UpstreamError('北交所分页总数变化')
                expected = total
            for row in content:
                if row.get('xxjsrq'):
                    report_dates.add(str(row['xxjsrq']))
                industry = next((row[k] for k in ('xxhyzl', 'sshymc', 'sshy', 'xxhy', 'industry') if row.get(k)), None)
                # xxsshy/所属行业 field aliases are explicit; no inference from company names.
                if industry is None:
                    industry = row.get('xxsshy') or row.get('所属行业')
                rows.append(security('BJ', row.get('xxzqdm') or row.get('zqdm'),
                                     row.get('xxzqjc') or row.get('zqjc'), industry,
                                     '北京证券交易所公开股票列表（原始报告日期另列；UTC发布时间未知）',
                                     'https://www.bse.cn/nq/listedcompany.html'))
            total_pages = data.get('totalPages')
            if (expected is not None and len(rows) >= expected) or (total_pages is not None and page + 1 >= int(total_pages)) or not content:
                break
        if expected is not None and expected != len(rows):
            raise UpstreamError('北交所分页覆盖不完整')
        self.reported['BJ'] = expected
        self.report_dates['BJ'] = sorted(report_dates)
        return rows
